from pathlib import Path

from docx import Document as DocxDocument

from rag.loaders.base import LoadedDoc


def load_docx(path: Path) -> LoadedDoc:
    doc = DocxDocument(str(path))
    lines = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    # 表格转 Markdown，避免表格语义丢失（rag_design#7）
    for table in doc.tables:
        rows = []
        for row in table.rows:
            rows.append("| " + " | ".join(cell.text.strip() for cell in row.cells) + " |")
        if rows:
            lines.append("\n".join(rows))
    return LoadedDoc(title=path.name, pages=[(0, "\n".join(lines))])
