from pathlib import Path

from openpyxl import load_workbook

from rag.loaders.base import LoadedDoc


def load_xlsx(path: Path) -> LoadedDoc:
    workbook = load_workbook(filename=str(path), data_only=True)
    pages: list[tuple[int, str]] = []
    for sheet_index, sheet in enumerate(workbook.worksheets, start=1):
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if cell is None else str(cell).strip() for cell in row]
            if any(cells):
                rows.append("| " + " | ".join(cells) + " |")
        if rows:
            pages.append((sheet_index, f"# {sheet.title}\n" + "\n".join(rows)))
    return LoadedDoc(title=path.name, pages=pages)
